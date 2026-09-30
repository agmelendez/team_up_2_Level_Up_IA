from docx import Document

PATH = 'Version 2.0/INA-Taller-IA-Idiomas-6oct2026.docx'

PARAGRAPH_REPLACEMENTS = {
    'El micro-entregable se deposita en el formulario único de la jornada al cierre de cada ventana; es el principal mecanismo pedagógico para asegurar la participación activa en un formato con cámaras cerradas.':
        'El producto de cada práctica se guarda en el cuaderno local del portal al cierre de cada ventana. Permanece en la computadora de la persona participante, puede descargarse como archivo Markdown y no se envía a ninguna plataforma. La participación activa se acompaña mediante pausas guiadas y devolución voluntaria en plenaria.',
    'Productos: número de respuestas recibidas por ciclo en el formulario único de la jornada, con una meta de participación de al menos el cuarenta por ciento en un ciclo o más, y catálogo de materiales devuelto al INA en formato editable dentro de los cinco días hábiles siguientes.':
        'Productos: cada participante conserva un cuaderno local acumulativo con sus instrucciones, evidencias, hallazgos críticos y próximos ajustes. No se recopilan ni contabilizan entregas individuales. El seguimiento durante la sesión se realiza mediante participación voluntaria y devolución en plenaria; el catálogo de materiales se entrega al INA en formato editable dentro de los cinco días hábiles siguientes.',
    'Deposite el resultado y su hallazgo principal en el formulario único de la jornada, o consulte el Plan B en el portal si requiere contrastar con una salida de respaldo.':
        'Guarde el resultado y su hallazgo principal en el cuaderno local del portal, o consulte el Plan B si requiere contrastar con una salida de respaldo. Descargue su copia antes de cerrar el navegador o cambiar de computadora.',
}

TABLE_REPLACEMENTS = {
    'Cada participante deposita su resultado en el formulario único de la jornada, con una sección por ciclo.':
        'Cada participante guarda su resultado en el cuaderno local del portal, con una sección por práctica y descarga personal en formato Markdown.',
    'Q&A + formulario único': 'Q&A + descarga del cuaderno local',
    'Microsoft Forms': 'Cuaderno local del portal',
    'Sondeos en vivo, formulario único de la jornada y encuestas de entrada y salida':
        'Registro personal de las siete prácticas y descarga del cuaderno; sondeos en vivo y encuestas institucionales solo si se habilitan por separado',
    'Configurar respuestas múltiples y sección por ciclo de práctica':
        'No requiere cuenta ni envío; verificar guardado local y descarga Markdown en el navegador de uso',
}

doc = Document(PATH)
for paragraph in doc.paragraphs:
    if paragraph.text in PARAGRAPH_REPLACEMENTS:
        paragraph.text = PARAGRAPH_REPLACEMENTS[paragraph.text]

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                original = paragraph.text
                updated = original
                for old, new in TABLE_REPLACEMENTS.items():
                    updated = updated.replace(old, new)
                if updated != original:
                    paragraph.text = updated

doc.save(PATH)
